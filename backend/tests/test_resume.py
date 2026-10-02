import io

import docx

from app.services import resume_gen, resume_parser
from .conftest import RESUME_TEXT, docx_bytes, pdf_bytes, register, upload


def test_docx_extraction_and_structure(client, auth):
    r = upload(client, auth)
    assert r.status_code == 201, r.text
    d = r.json()
    assert d["processing_status"] == "parsed" and d["file_type"] == "docx"
    p = d["structured_profile"]
    assert p["name"] == "Jane Doe" and p["email"] == "jane@example.com"
    assert {"Python", "SQL", "PostgreSQL", "Docker"} <= set(p["skills"])
    e = p["experience"]
    assert len(e) == 2
    assert (e[0]["title"], e[0]["company"], e[0]["start"], e[0]["end"]) == ("Senior Software Engineer", "Initech", "2020-01", "present")
    assert (e[1]["title"], e[1]["company"]) == ("Software Developer", "Hooli")
    assert len(e[0]["bullets"]) == 3 and "2 million" in e[0]["bullets"][0]
    assert p["certifications"] == ["AWS Certified Developer"]
    assert p["education"][0]["year"] == "2015"


def test_pdf_extraction(client, auth):
    r = upload(client, auth, pdf_bytes(), "resume.pdf")
    assert r.status_code == 201, r.text
    assert "Jane Doe" in r.json()["extracted_text"]
    assert "Python" in r.json()["structured_profile"]["skills"]


def test_rejects_bad_files(client, auth):
    assert upload(client, auth, b"hello", "resume.txt").status_code == 415
    assert upload(client, auth, b"", "resume.pdf").status_code in (415, 422)  # multipart drops empty parts
    assert upload(client, auth, b"not a pdf at all", "resume.pdf").status_code == 415  # magic bytes mismatch
    assert upload(client, auth, b"%PDF-1.4 corrupted garbage", "resume.pdf").status_code == 422
    assert upload(client, auth, b"PK\x03\x04 garbage", "resume.docx").status_code == 422
    empty = docx.Document()
    b = io.BytesIO()
    empty.save(b)
    r = upload(client, auth, b.getvalue(), "empty.docx")
    assert r.status_code == 422 and "No extractable text" in r.text


def test_encrypted_pdf_rejected(client, auth):
    from pypdf import PdfReader, PdfWriter
    w = PdfWriter()
    for p in PdfReader(io.BytesIO(pdf_bytes())).pages:
        w.add_page(p)
    w.encrypt("secret")
    b = io.BytesIO()
    w.write(b)
    r = upload(client, auth, b.getvalue(), "locked.pdf")
    assert r.status_code == 422 and "password" in r.text.lower()


def test_oversize_rejected(client, auth, monkeypatch):
    from app.config import get_settings
    monkeypatch.setattr(get_settings(), "max_upload_mb", 0)
    assert upload(client, auth).status_code == 413


def test_duplicate_upload_and_original_preserved(client, auth):
    r = upload(client, auth)
    assert upload(client, auth).status_code == 409
    link = client.get(f"/api/resumes/{r.json()['id']}/download-link", headers=auth).json()["url"]
    body = client.get(link)
    assert body.status_code == 200 and body.content == docx_bytes()  # byte-identical original


def test_user_can_correct_parsed_data(client, auth):
    rid = upload(client, auth).json()["id"]
    sp = client.get(f"/api/resumes/{rid}", headers=auth).json()["structured_profile"]
    sp["skills"].append("FastAPI")
    r = client.patch(f"/api/resumes/{rid}", headers=auth, json={"structured_profile": sp})
    assert "FastAPI" in r.json()["structured_profile"]["skills"]


# ---------------- generation + factual integrity ----------------
SRC = resume_parser.parse_structured(RESUME_TEXT)
JOB = "We need a Senior Software Engineer. Required: Python, SQL, Kubernetes, Terraform. Mentoring is a plus. 5+ years."


def test_generation_only_reorders():
    content, analysis = resume_gen.generate(SRC, "Senior Software Engineer", JOB)
    assert sorted(content["skills"]) == sorted(SRC["skills"])
    assert content["skills"][:2] == ["Python", "SQL"]  # job-relevant skills first
    assert resume_gen.validate(content, SRC, JOB)["passed"]
    assert "Kubernetes" in analysis["missing_requirements"] and "Terraform" in analysis["missing_requirements"]
    assert "Kubernetes" not in content["skills"]  # gap is reported, never added
    assert analysis["keyword_coverage"] is not None


def test_validator_flags_every_kind_of_invention():
    content, _ = resume_gen.generate(SRC, "Senior Software Engineer", JOB)
    bad = {**content, "skills": content["skills"] + ["Kubernetes"]}
    codes = lambda c: {i["code"] for i in resume_gen.validate(c, SRC, JOB)["issues"] if i["severity"] == "error"}  # noqa: E731
    assert "unsupported_skill" in codes(bad)

    bad = {**content, "experience": [{**content["experience"][0], "title": "Principal Engineer"}, *content["experience"][1:]]}
    assert {"unsupported_role", "role_omitted"} <= codes(bad)

    e0 = content["experience"][0]
    bad = {**content, "experience": [{**e0, "bullets": e0["bullets"] + ["Reduced costs by 40% using Kubernetes"]}, *content["experience"][1:]]}
    assert {"unsupported_number", "unsupported_term"} <= codes(bad) | {"unsupported_term"}
    assert "unsupported_number" in codes(bad)

    assert "unsupported_certification" in codes({**content, "certifications": ["CISSP"]})
    assert "unsupported_education" in codes({**content, "education": [{"text": "PhD, MIT"}]})
    assert "role_omitted" in codes({**content, "experience": content["experience"][:1]})
    assert "hidden_text" in codes({**content, "summary": "Backend​ engineer"})


def test_job_skill_not_claimed_via_summary():
    content, _ = resume_gen.generate(SRC, "Senior Software Engineer", JOB)
    c = {**content, "summary": "Backend engineer experienced with Terraform"}
    r = resume_gen.validate(c, SRC, JOB)
    assert not r["passed"] and any("terraform" in i["detail"].lower() for i in r["issues"])


def test_ats_docx_output_is_simple_and_complete():
    content, _ = resume_gen.generate(SRC, "Senior Software Engineer", JOB)
    d = docx.Document(io.BytesIO(resume_gen.render_docx(content)))
    heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    assert {"Skills", "Experience", "Education"} <= set(heads)
    text = "\n".join(p.text for p in d.paragraphs)
    assert "Initech" in text and "2020-01 – present" in text
    assert not d.tables and not d.inline_shapes  # no tables/images that trip ATS parsers
    assert all(not s.header.paragraphs[0].text for s in d.sections)  # nothing hidden in headers
    assert resume_gen.render_pdf(content).startswith(b"%PDF")


def test_cover_letter_uses_only_documented_facts():
    txt = resume_gen.cover_letter(SRC, "Senior Software Engineer", "Globex", ["Python", "SQL"])
    assert resume_gen.validate_text(txt, SRC)["passed"]
    assert not resume_gen.validate_text(txt + " I led a Kubernetes migration saving 30%.", SRC)["passed"]


def test_version_history_and_edit_creates_new_version(client, auth, web):
    from .conftest import gh_job, make_profile, run_now
    rid = upload(client, auth).json()["id"]
    web.greenhouse["acme"] = [gh_job(1, "Senior Software Engineer")]
    pid = make_profile(client, auth)["id"]
    run_now(pid)
    job_id = client.get("/api/jobs", headers=auth).json()["items"][0]["id"]
    client.post(f"/api/jobs/{job_id}/prepare-application", headers=auth, json={})
    vs = client.get(f"/api/resumes/{rid}/versions", headers=auth).json()
    assert len(vs) == 1 and vs[0]["version_number"] == 1 and vs[0]["validation_results"]["passed"]
    content = vs[0]["generated_content"]
    content["skills"].append("Kubernetes")
    e = client.post(f"/api/resume-versions/{vs[0]['id']}/edit", headers=auth, json={"content": content}).json()
    assert e["version_number"] == 2 and not e["validation_results"]["passed"]
    assert len(client.get(f"/api/resumes/{rid}/versions", headers=auth).json()) == 2  # original untouched
