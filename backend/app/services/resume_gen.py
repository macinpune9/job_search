"""ATS-friendly resume generation and factual-integrity validation.

The generator is deterministic: it can only SELECT and REORDER facts that already exist in the candidate's
structured profile. It never adds, rewrites or infers content. The validator is independent of the generator
and is also run on user-edited versions and on any future LLM-produced output.
"""
from __future__ import annotations

import io
import re

from .textutil import contains_term, fold, tokens

SKILL_VOCAB = """python java javascript typescript golang rust c++ c# ruby php swift kotlin scala sql nosql postgresql mysql
mongodb redis elasticsearch kafka rabbitmq spark hadoop airflow dbt snowflake bigquery redshift aws azure gcp kubernetes docker
terraform ansible jenkins gitlab github ci/cd devops linux react angular vue node.js django flask fastapi spring .net graphql
grpc microservices machine learning deep learning nlp pytorch tensorflow scikit-learn pandas numpy tableau power bi excel
looker salesforce sap oracle jira agile scrum kanban six sigma pmp prince2 seo sem google analytics figma sketch photoshop
iso 27001 gdpr hipaa soc 2 pci selenium cypress junit pytest cucumber postman api design system design data modeling etl
leadership stakeholder management budgeting forecasting p&l negotiation recruiting coaching mentoring cissp cisa cpa cfa itil
""".split("\n")
_MULTI = ["machine learning", "deep learning", "power bi", "google analytics", "six sigma", "api design", "system design",
          "data modeling", "stakeholder management", "iso 27001", "soc 2"]
VOCAB: list[str] = sorted({w for line in SKILL_VOCAB for w in line.split() if w not in
                           {"machine", "learning", "deep", "power", "bi", "google", "analytics", "six", "sigma", "api", "design",
                            "system", "data", "modeling", "stakeholder", "management", "iso", "27001", "soc", "2", "ci/cd"}}
                          | set(_MULTI) | {"ci/cd"})
STOP = set("""the and for with that this from have has are was were will you your our their they them but not all any can
into over under about than then also such more most other some what when where which who whom while within without per via
etc able work works working team teams role roles job company years year new use used using ability strong experience
skills including include includes responsible responsibilities required preferred plus must should may across each both""".split())


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", fold(s or "").lower()).strip()


def _flat_text(c: dict) -> str:
    parts = [c.get("name") or "", c.get("summary") or "", " ".join(c.get("skills", []))]
    for e in c.get("experience", []):
        parts += [e.get("title", ""), e.get("company", ""), " ".join(e.get("bullets", []))]
    parts += [x.get("text") or "" for x in c.get("education", [])]
    parts += c.get("certifications", []) + c.get("achievements", [])
    return "\n".join(parts)


def job_terms(job_text: str, candidate_terms: list[str]) -> list[str]:
    """Skill-like terms present in the job text (controlled vocabulary + the candidate's own skills)."""
    seen, out = set(), []
    for t in list(VOCAB) + candidate_terms:
        if t.lower() not in seen and contains_term(job_text, t):
            seen.add(t.lower())
            m = re.search(r"(?<![A-Za-z0-9+#])" + re.escape(t) + r"(?![A-Za-z0-9+#])", job_text, re.I)
            out.append(m.group(0) if m else t)  # keep the job posting's own capitalisation
    return out


def generate(structured: dict, job_title: str, job_text: str, company: str = "") -> tuple[dict, dict]:
    """Return (content, analysis). Only reorders existing facts. `company` is excluded from gap detection."""
    jt = tokens(f"{job_title} {job_text}")
    src_skills = list(structured.get("skills", []))
    in_job = [s for s in src_skills if contains_term(job_text, s)]
    rest = [s for s in src_skills if s not in in_job]
    content = {k: (list(v) if isinstance(v, list) else v) for k, v in structured.items()}
    content["skills"] = in_job + rest
    changes = []
    if in_job and content["skills"] != src_skills:
        changes.append({"type": "reorder_skills", "detail": "Skills mentioned in the job description moved first.",
                        "reason": "Improves relevance for ATS/keyword screens; no skills added or removed."})

    def rel(b: str) -> int:
        return len(tokens(b) & jt - STOP)

    exp = []
    for e in structured.get("experience", []):
        b = list(e.get("bullets", []))
        nb = sorted(b, key=lambda x: -rel(x))  # stable
        if nb != b:
            changes.append({"type": "reorder_bullets", "detail": f"{e.get('title')} @ {e.get('company')}",
                            "reason": "Bullets most relevant to the job moved up; wording unchanged."})
        exp.append({**e, "bullets": nb})
    content["experience"] = exp

    own = {w for w in tokens(company) if len(w) > 2}
    all_terms = [t for t in job_terms(job_text, src_skills) if t.lower() not in own]  # the employer's name is not a skill
    src_text = _flat_text(structured)
    matched = [t for t in all_terms if contains_term(src_text, t)]
    missing = [t for t in all_terms if not contains_term(src_text, t)]
    analysis = {
        "matched_skills": matched, "missing_requirements": missing,
        "keyword_coverage": round(len(matched) / len(all_terms), 2) if all_terms else None,
        "changes": changes,
        "note": "Missing items are gaps: they are NOT added to the resume.",
    }
    return content, analysis


def validate(content: dict, source: dict, job_text: str = "") -> dict:
    """Check `content` introduces no facts absent from `source`. Returns {passed, issues}."""
    issues: list[dict] = []

    def bad(code, detail):
        issues.append({"code": code, "detail": detail, "severity": "error"})

    src_text = _flat_text(source)
    cont_text = _flat_text(content)
    src_norm = _norm(src_text)

    src_sk = {_norm(s) for s in source.get("skills", [])}
    for s in content.get("skills", []):
        if _norm(s) not in src_sk and _norm(s) not in src_norm:
            bad("unsupported_skill", s)

    def key(e):
        return (_norm(e.get("title", "")), _norm(e.get("company", "")), e.get("start"), e.get("end"))

    src_keys = [key(e) for e in source.get("experience", [])]
    got_keys = [key(e) for e in content.get("experience", [])]
    for k in got_keys:
        if k not in src_keys:
            bad("unsupported_role", f"{k[0]} @ {k[1]} ({k[2]}–{k[3]}) is not in the source resume")
    for k in src_keys:
        if k not in got_keys:
            bad("role_omitted", f"{k[0]} @ {k[1]} missing; employment history must stay complete")

    src_bul = {_norm(b) for e in source.get("experience", []) for b in e.get("bullets", [])}
    for e in content.get("experience", []):
        for b in e.get("bullets", []):
            if _norm(b) not in src_bul:
                # edited wording is allowed only if it adds no new numbers or job-only/skill terms (checked below)
                issues.append({"code": "bullet_reworded", "detail": b[:120], "severity": "warning"})

    for field, label in (("certifications", "certification"), ("achievements", "achievement")):
        have = {_norm(x) for x in source.get(field, [])}
        for x in content.get(field, []):
            if _norm(x) not in have:
                bad(f"unsupported_{label}", x)
    have_edu = {_norm(x.get("text") or "") for x in source.get("education", [])}
    for x in content.get("education", []):
        if _norm(x.get("text") or "") not in have_edu:
            bad("unsupported_education", x.get("text") or "")

    if _norm(content.get("summary") or "") not in ("", _norm(source.get("summary") or "")):
        issues.append({"code": "summary_changed", "detail": "Summary differs from source; verified below.", "severity": "warning"})

    for n in set(re.findall(r"\d[\d,.]*%?", cont_text)) - set(re.findall(r"\d[\d,.]*%?", src_text)):
        bad("unsupported_number", f"'{n}' does not appear in the source resume (possible invented metric)")

    cont_tok, src_tok = tokens(cont_text), tokens(src_text)
    for t in sorted((cont_tok - src_tok) & (tokens(job_text) - STOP)):
        if len(t) > 2:
            bad("job_term_not_in_resume", f"'{t}' appears in the job description but not in the source resume")
    for t in VOCAB:
        if contains_term(cont_text, t) and not contains_term(src_text, t):
            if not any(i["code"] == "job_term_not_in_resume" and t in i["detail"] for i in issues):
                bad("unsupported_term", t)

    if re.search(r"[​-‏⁠﻿]", cont_text):
        bad("hidden_text", "zero-width/invisible characters found")

    # dedupe
    seen, uniq = set(), []
    for i in issues:
        k = (i["code"], i["detail"])
        if k not in seen:
            seen.add(k)
            uniq.append(i)
    return {"passed": not any(i["severity"] == "error" for i in uniq), "issues": uniq}


# ----------------------------------------------------------------------------------------
def render_docx(c: dict) -> bytes:
    import docx
    from docx.shared import Pt

    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    for h in ("Heading 1", "Heading 2"):
        d.styles[h].font.name = "Calibri"
    d.add_heading(c.get("name") or "Candidate", 0)
    contact = " | ".join(x for x in (c.get("email"), c.get("phone"), c.get("location")) if x)
    if contact:
        d.add_paragraph(contact)
    if c.get("summary"):
        d.add_heading("Summary", 1)
        d.add_paragraph(c["summary"])
    if c.get("skills"):
        d.add_heading("Skills", 1)
        d.add_paragraph(", ".join(c["skills"]))
    if c.get("experience"):
        d.add_heading("Experience", 1)
        for e in c["experience"]:
            d.add_heading(f"{e.get('title', '')}, {e.get('company', '')}".strip(", "), 2)
            d.add_paragraph(f"{e.get('start', '')} – {e.get('end', '')}")
            for b in e.get("bullets", []):
                d.add_paragraph(b, style="List Bullet")
    if c.get("education"):
        d.add_heading("Education", 1)
        for x in c["education"]:
            d.add_paragraph(x.get("text", ""))
    if c.get("certifications"):
        d.add_heading("Certifications", 1)
        for x in c["certifications"]:
            d.add_paragraph(x, style="List Bullet")
    if c.get("achievements"):
        d.add_heading("Achievements", 1)
        for x in c["achievements"]:
            d.add_paragraph(x, style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def render_pdf(c: dict) -> bytes:
    from xml.sax.saxutils import escape

    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import ListFlowable, Paragraph, SimpleDocTemplate, Spacer

    ss = getSampleStyleSheet()
    story = [Paragraph(escape(c.get("name") or "Candidate"), ss["Title"])]
    contact = " | ".join(x for x in (c.get("email"), c.get("phone"), c.get("location")) if x)
    if contact:
        story.append(Paragraph(escape(contact), ss["Normal"]))

    def sec(title, text=None, bullets=None):
        story.append(Spacer(1, 8))
        story.append(Paragraph(title, ss["Heading2"]))
        if text:
            story.append(Paragraph(escape(text), ss["Normal"]))
        if bullets:
            story.append(ListFlowable([Paragraph(escape(b), ss["Normal"]) for b in bullets], bulletType="bullet"))

    if c.get("summary"):
        sec("Summary", c["summary"])
    if c.get("skills"):
        sec("Skills", ", ".join(c["skills"]))
    if c.get("experience"):
        story.append(Spacer(1, 8))
        story.append(Paragraph("Experience", ss["Heading2"]))
        for e in c["experience"]:
            story.append(Paragraph(escape(f"{e.get('title', '')}, {e.get('company', '')}"), ss["Heading3"]))
            story.append(Paragraph(escape(f"{e.get('start', '')} – {e.get('end', '')}"), ss["Normal"]))
            if e.get("bullets"):
                story.append(ListFlowable([Paragraph(escape(b), ss["Normal"]) for b in e["bullets"]], bulletType="bullet"))
    if c.get("education"):
        sec("Education", None, [x.get("text", "") for x in c["education"]])
    if c.get("certifications"):
        sec("Certifications", None, c["certifications"])
    if c.get("achievements"):
        sec("Achievements", None, c["achievements"])
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4).build(story)
    return buf.getvalue()


def cover_letter(structured: dict, job_title: str, company: str, matched: list[str]) -> str:
    """Template letter built only from documented facts. Review and edit before use."""
    name = structured.get("name") or "Candidate"
    exp = (structured.get("experience") or [None])[0]
    lines = [f"Dear {company} hiring team,", "",
             f"I am writing to apply for the {job_title} position at {company}."]
    if exp and exp.get("title"):
        lines.append(f"My most recent role is {exp['title']}" + (f" at {exp['company']}." if exp.get("company") else "."))
    if matched:
        lines.append("My documented experience includes " + ", ".join(matched[:8]) + ", which overlaps with the requirements in your posting.")
    lines += ["", "Thank you for your consideration.", "", f"Sincerely,", name]
    return "\n".join(lines)


def validate_text(text: str, source: dict) -> dict:
    """Lighter check for free text (cover letters): no new numbers or skill/tool terms absent from the resume."""
    src_text = _flat_text(source)
    issues = []
    for n in set(re.findall(r"\d[\d,.]*%?", text)) - set(re.findall(r"\d[\d,.]*%?", src_text)):
        issues.append({"code": "unsupported_number", "detail": n, "severity": "error"})
    for t in VOCAB:
        if contains_term(text, t) and not contains_term(src_text, t):
            issues.append({"code": "unsupported_term", "detail": t, "severity": "error"})
    return {"passed": not issues, "issues": issues}
