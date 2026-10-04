"""Resume text extraction and heuristic structuring.

Extraction is deterministic. Structuring is rule-based and best-effort: the user is expected to review
and correct it in the Resume Manager (the corrected profile is the source of truth for generation).
"""
from __future__ import annotations

import io
import re
import zipfile

from .skillvocab import is_plausible_skill

MAGIC = {"pdf": b"%PDF-", "docx": b"PK"}
MAX_UNCOMPRESSED = 50 * 1024 * 1024


class ResumeError(Exception):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


def detect_type(filename: str, data: bytes) -> str:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext not in MAGIC:
        raise ResumeError("unsupported_type", "Only PDF and DOCX files are supported.")
    if not data:
        raise ResumeError("empty_file", "The file is empty.")
    if not data.startswith(MAGIC[ext]):
        raise ResumeError("unsupported_type", "File content does not match its extension.")
    return ext


def extract_text(file_type: str, data: bytes) -> str:
    text = _pdf(data) if file_type == "pdf" else _docx(data)
    text = text.replace("\x00", "").strip()
    if len(text) < 20:
        raise ResumeError("empty_document", "No extractable text found (scanned image PDFs are not supported).")
    return text


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PyPdfError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ResumeError("encrypted", "The PDF is password-protected. Remove the password and re-upload.")
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    except ResumeError:
        raise
    except (PyPdfError, ValueError, KeyError, OSError, AttributeError, TypeError) as e:
        raise ResumeError("corrupted", f"The PDF could not be read ({type(e).__name__}).")


def _docx(data: bytes) -> str:
    import docx

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:  # zip-bomb guard
            if sum(i.file_size for i in z.infolist()) > MAX_UNCOMPRESSED:
                raise ResumeError("corrupted", "Document expands to an unsafe size.")
            names = z.namelist()
            if any(n.startswith("EncryptedPackage") or n == "EncryptionInfo" for n in names):
                raise ResumeError("encrypted", "The document is encrypted.")
        d = docx.Document(io.BytesIO(data))
    except ResumeError:
        raise
    except Exception as e:  # noqa: BLE001 - python-docx raises many types for bad input
        raise ResumeError("corrupted", f"The DOCX could not be read ({type(e).__name__}).")
    lines = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(lines)


# ----------------------------------------------------------------------------------------
MONTHS = {m: i + 1 for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"])}
_D = r"(?:(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?,?\s+\d{4}|\d{1,2}/\d{4}|\d{4})"
DATE_RANGE = re.compile(rf"(?P<s>{_D})\s*(?:-|–|—|to)\s*(?P<e>{_D}|present|current|now|today)", re.I)
HEADINGS = {
    "summary": r"(professional |career |executive )?(profile )?(summary|profile|objective)|profile summary|about( me)?",
    "experience": r"(professional |work |relevant |project )?(experience|employment( history)?|work history|career history)|project (details|history|summary)|professional background",
    "education": r"education( and training)?|academic.*|education(al)? (background|qualifications?)|qualifications",
    "skills": r"(technical |key |core )?(skills|competenc(y|ies)|technologies|tech stack|tools)(.*)?",
    "certifications": r"(licenses? (and|&) )?certifications?( (and|&) (trainings?|courses|licenses?))?|licenses?|courses|trainings?( (and|&) certifications?)?",
    "achievements": r"(key )?(achievements|awards|accomplishments|honou?rs)",
    "projects": r"(personal |key |selected )?projects?",
    # sections that are NOT skills: recognising them ends the Skills section so their lines never leak into it
    "interests": r"(hobbies|interests|hobbies (and|&) interests|personal interests|activities|extracurricular( activities)?|leisure)",
    "languages": r"(spoken |foreign )?languages?( known| spoken)?( skills| proficiency)?|sprachen|langues",
    "personal": r"(personal|contact)( (details|information|data|profile))?|references?|availability|nationality|work (permit|authori[sz]ation)|bildung|persönliche daten|coordonn[ée]es",
    "volunteer": r"volunteer(ing)?( experience| work)?|community( service)?|publications?|memberships?|workshops?",
}
# "Testing Tools JIRA with Zephyr" -> "JIRA with Zephyr", "Programming Languages Java" -> "Java" (a short category label + more text)
_SKILL_LABEL = re.compile(r"^(?:[A-Za-z/&+\-]+\s+){0,2}(?:tools?|languages?|frameworks?|technologies)\s+(?=\S)", re.I)
BULLET = re.compile(r"^\s*([-•*·▪●◦–]|\d+\.)\s+")
_COMPANY = re.compile(r"^\s*(?:company|client|employer|organi[sz]ation)\s*[:\-–]?\s+(\S.*)$", re.I)
_ROLE = re.compile(r"^\s*(?:role|position|designation|job title|title)\s*[:\-–]?\s+(\S.*)$", re.I)
_PERIOD_LABEL = re.compile(r"^(period|duration|dates?|tenure|from)\s*[:\-–]?\s*", re.I)
_EDU_LINE = re.compile(r"(universit|college|school|bachelor|master|degree|diploma|b\.?sc|m\.?sc|certif|course|training|gymnasium|hochschule)", re.I)


def _labels(lines):
    """(title, company) from labelled lines such as "Role: QA Engineer" / "Company Acme AG" (consultant-style CVs)."""
    title = company = None
    for l in lines:
        m = _ROLE.match(l or "")
        if m and not title:
            title = m.group(1).strip(" .,")
        m = _COMPANY.match(l or "")
        if m and not company:
            company = m.group(1).strip(" .,")
    return title, company


def _implicit_experience(lines, section_at):
    """CVs without an Experience heading: build one entry per date range found outside education/skills/etc."""
    ok = lambda k: section_at[k] in (None, "experience", "projects")  # noqa: E731
    idx = [i for i, l in enumerate(lines) if DATE_RANGE.search(l) and ok(i) and not _EDU_LINE.search(l) and not (i and _EDU_LINE.search(lines[i - 1]))]
    out = []
    for n, i in enumerate(idx):
        nxt = idx[n + 1] if n + 1 < len(idx) else len(lines)
        lo = max(idx[n - 1] + 1 if n else 0, i - 8)
        pre = [lines[k] for k in range(lo, i) if lines[k].strip() and ok(k)]
        post = [lines[k] for k in range(i + 1, nxt) if lines[k].strip() and ok(k)]
        cut = next((k for k in range(len(post) - 1, max(-1, len(post) - 9), -1) if _COMPANY.match(post[k])), None)
        if cut is not None and n + 1 < len(idx):
            post = post[:cut]  # these lines are the NEXT block's header
        dm = DATE_RANGE.search(lines[i])
        rest = _PERIOD_LABEL.sub("", (lines[i][:dm.start()] + " " + lines[i][dm.end():]).strip(" |,-–—()")).strip(" |,-–—()")
        title, company = _labels([*pre, *post[:3]])
        if not (title or company):
            title, company = _split_header(pre[-2:] + ([rest] if rest else []))
        out.append({"title": title or "", "company": company or "", "start": _norm_date(dm.group("s")), "end": _norm_date(dm.group("e")),
                    "bullets": [_bullet_text(x) for x in post if not _ROLE.match(x)][:40]})
    return out
EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE = re.compile(r"(?<!\d)(\+?\d[\d\s().-]{7,}\d)")
DEGREE = re.compile(r"\b(bachelor|master|b\.?sc|m\.?sc|b\.?a\b|m\.?a\b|b\.?eng|m\.?eng|b\.?tech|m\.?tech|ph\.?d|doctorate|mba|diploma|associate)", re.I)


def _norm_date(s: str) -> str:
    s = s.strip().lower().rstrip(",")
    if s in ("present", "current", "now", "today"):
        return "present"
    m = re.match(r"([a-z]{3})[a-z]*\.?,?\s+(\d{4})", s)
    if m:
        return f"{m.group(2)}-{MONTHS[m.group(1)]:02d}"
    m = re.match(r"(\d{1,2})/(\d{4})", s)
    if m:
        return f"{m.group(2)}-{int(m.group(1)):02d}"
    return s


def _heading(line: str) -> str | None:
    t = line.strip().strip(":").strip()
    if not t or len(t) > 40:
        return None
    for k, pat in HEADINGS.items():
        if re.fullmatch(pat, t, re.I):
            return k
    return None


def _split_header(lines: list[str]) -> tuple[str, str]:
    lines = [l.strip(" ,|-–—") for l in lines if l.strip()]
    if not lines:
        return "", ""
    if len(lines) >= 2:
        return _order(lines[-2], lines[-1])
    h = lines[0]
    for sep in (" at ", " @ ", " | ", " – ", " — ", " - ", ", "):
        if sep in h:
            a, b = h.split(sep, 1)
            return _order(a.strip(), b.strip())
    return h, ""


_TITLE_WORDS = re.compile(r"\b(engineer|developer|manager|analyst|director|lead|consultant|specialist|designer|scientist|"
                          r"architect|administrator|officer|intern|associate|coordinator|head|vp|president|technician)\b", re.I)


def _order(a: str, b: str) -> tuple[str, str]:
    """Return (title, company); swap when only the second part looks like a job title."""
    if _TITLE_WORDS.search(b) and not _TITLE_WORDS.search(a):
        return b, a
    return a, b


def _bullet_text(s: str) -> str:
    return BULLET.sub("", s).strip()


def parse_structured(text: str) -> dict:
    lines = [l.rstrip() for l in text.splitlines()]
    nonempty = [l.strip() for l in lines if l.strip()]
    out: dict = {"name": None, "email": None, "phone": None, "summary": "", "skills": [], "experience": [],
                 "education": [], "certifications": [], "achievements": []}
    if nonempty and "@" not in nonempty[0] and not re.search(r"\d", nonempty[0]) and len(nonempty[0].split()) <= 5 \
            and not _heading(nonempty[0]):
        out["name"] = nonempty[0]
    m = EMAIL.search(text)
    out["email"] = m.group(0) if m else None
    head = "\n".join(nonempty[:8])
    m = PHONE.search(head)
    out["phone"] = m.group(1).strip() if m else None

    sections: dict[str, list[str]] = {}
    section_at: list = []
    cur = None
    for l in lines:
        h = _heading(l)
        if h:
            cur = h
            sections.setdefault(cur, [])
        elif cur and l.strip():
            sections[cur].append(l)
        section_at.append(cur)

    out["summary"] = " ".join(x.strip() for x in sections.get("summary", []))[:1500]

    skills: list[str] = []
    for l in sections.get("skills", []):
        l = _bullet_text(l)
        if ":" in l and len(l.split(":", 1)[0]) < 30:
            l = l.split(":", 1)[1]
        for p in re.split(r"[,;|•·]|\.\s+", l):
            p = _SKILL_LABEL.sub("", p.strip(" ."), count=1).strip(" .")
            if p and len(p) <= 40 and p.lower() not in {s.lower() for s in skills}:
                skills.append(p)
    out["skills"] = [s for s in skills if is_plausible_skill(s)]

    exp: list[dict] = []
    pending: list[str] = []
    cur_e: dict | None = None
    cur_bul: list[tuple[str, bool]] = []

    def close():
        if cur_e is not None:
            cur_e["bullets"] = [t for t, _ in cur_bul]

    for l in sections.get("experience", []):
        dm = DATE_RANGE.search(l)
        if dm:
            # unmarked short lines trailing the previous entry belong to this entry's header
            hdr_prev: list[str] = []
            lab_at = None
            if cur_e is not None:  # a labelled header block ("Company ...") among the last lines belongs to THIS entry
                lab_at = next((k for k in range(len(cur_bul) - 1, max(-1, len(cur_bul) - 9), -1) if _COMPANY.match(cur_bul[k][0]) and not cur_bul[k][1]), None)
            if lab_at is not None:
                hdr_prev = [t for t, _ in cur_bul[lab_at:]]
                del cur_bul[lab_at:]
            else:
                while cur_e is not None and cur_bul and not cur_bul[-1][1] and len(cur_bul[-1][0]) < 80 and len(hdr_prev) < 2:
                    hdr_prev.insert(0, cur_bul.pop()[0])
            close()
            rest = _PERIOD_LABEL.sub("", (l[:dm.start()] + " " + l[dm.end():]).strip(" |,-–—()")).strip(" |,-–—()")
            lab_title, lab_company = _labels([*hdr_prev, *pending, rest])
            if lab_title or lab_company:
                title, company = lab_title or "", lab_company or ""
            elif rest:
                title, company = _split_header([*hdr_prev, rest])
            else:
                title, company = _split_header(hdr_prev or pending)
            cur_e = {"title": title, "company": company, "start": _norm_date(dm.group("s")),
                     "end": _norm_date(dm.group("e")), "bullets": []}
            exp.append(cur_e)
            cur_bul = []
            pending = []
        elif cur_e is None:
            pending.append(l.strip())
        else:
            rm = _ROLE.match(l)
            if rm and not cur_e["title"]:
                cur_e["title"] = rm.group(1).strip(" .,")
            else:
                cur_bul.append((_bullet_text(l), bool(BULLET.match(l))))
    close()
    out["experience"] = exp or _implicit_experience(lines, section_at)
    out["parsed_by"] = "rules"

    for l in sections.get("education", []):
        l = _bullet_text(l)
        ym = re.search(r"(19|20)\d{2}", l)
        out["education"].append({"text": l, "degree": l if DEGREE.search(l) else None, "year": ym.group(0) if ym else None})
    out["certifications"] = [_bullet_text(l) for l in sections.get("certifications", [])]
    out["achievements"] = [_bullet_text(l) for l in sections.get("achievements", [])]
    return out

